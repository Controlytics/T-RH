"""Builds Dashboard-Screens.pptx from dashboard-shots/*.png.

2026-10-03, operator: "make a presentation with all dashboard screens proper
good presentation ppt" / "just with images". One full-bleed screenshot per
slide with a small screen-name label, a cover and a closing slide. No prose,
by instruction.

    python Layouts/presentation/build_dashboard_deck.py

The screenshots are real screens of the local app at 1600x900, captured on
2026-10-03 through a throwaway HTTP dev server (see CHANGELOG).
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
SHOTS = HERE / "dashboard-shots"
OUT = HERE / "Dashboard-Screens.pptx"

# Order and label of every slide. The label is the only text on a screen slide.
SLIDES = [
    ("01-login.png", "Sign in"),
    ("02-dashboard-overview.png", "Dashboard — overview"),
    ("03-block-areas.png", "Dashboard — CWH Block areas"),
    ("04-all-instruments-tabular.png", "All Instruments — tabular view"),
    ("05-area-alarms-panel.png", "Area — alarms in this selection"),
    ("06-qc-graphical.png", "QC — graphical view"),
    ("07-qc-fullscreen.png", "QC — floor plan, full screen"),
    ("08-microbiology-fullscreen.png", "QC Microbiology — floor plan"),
    ("09-qc3-fullscreen.png", "QC 3 — floor plan"),
    ("10-module-e-fullscreen.png", "Module-E — floor plan"),
    ("11-module-f-ground-fullscreen.png", "Module-F — ground floor"),
    ("12-module-f-first-fullscreen.png", "Module-F — first floor"),
    ("13-stability-fullscreen.png", "Stability — floor plan"),
    ("14-trend-page.png", "Instrument trend"),
    ("15-trend-charts.png", "Instrument trend — charts"),
    ("16-alarms.png", "Alarms"),
    ("17-data-report.png", "Data Report"),
    ("18-settings.png", "Settings"),
    ("19-reports.png", "Reports"),
    ("20-instruments.png", "Instruments"),
    ("21-functional-locations.png", "Functional Locations"),
    ("22-audit-trail.png", "Audit Trail"),
    ("23-configuration.png", "Configuration"),
]

NAVY = RGBColor(0x0F, 0x1E, 0x3A)
INK = RGBColor(0x0F, 0x17, 0x2A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0xB8, 0xC4, 0xD6)
ACCENT = RGBColor(0x5B, 0x9C, 0xF0)

W = Inches(13.333)
H = Inches(7.5)


def fill(slide, color):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = color


def circle(slide, x, y, d, alpha):
    """A soft white circle — the app's own band decoration."""
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, d, d)
    s.fill.solid()
    s.fill.fore_color.rgb = WHITE
    s.line.fill.background()
    # python-pptx has no transparency API; write the alpha into the XML.
    from pptx.oxml.ns import qn
    from lxml import etree
    srgb = s.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    a = etree.SubElement(srgb, qn("a:alpha"))
    a.set("val", str(int(alpha * 1000)))  # percent * 1000
    return s


def text(slide, x, y, w, h, s, size, bold=False, color=WHITE, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = s
    r.font.name = "Calibri"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return tb


def main():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    # Cover
    s = prs.slides.add_slide(blank)
    fill(s, NAVY)
    circle(s, Inches(9.4), Inches(-2.2), Inches(6), 10)
    circle(s, Inches(-1.6), Inches(4.6), Inches(4.6), 7)
    text(s, Inches(0.9), Inches(2.35), Inches(11.5), Inches(1.2), "T&RH-Monitoring", 54, bold=True)
    text(s, Inches(0.9), Inches(3.55), Inches(11.5), Inches(0.8), "Dashboard screens", 28, color=MUTED)
    text(s, Inches(0.9), Inches(6.5), Inches(11.5), Inches(0.4), "Controlytics AI Pvt Ltd  ·  21 CFR Part 11", 14, color=MUTED)

    # One screen per slide
    for file, label in SLIDES:
        p = SHOTS / file
        if not p.exists():
            raise SystemExit(f"missing screenshot: {p}")
        s = prs.slides.add_slide(blank)
        fill(s, INK)
        # Full bleed: the capture is 16:9 and so is the slide.
        s.shapes.add_picture(str(p), 0, 0, W, H)
        # The one word on the slide: which screen this is, in a small dark
        # pill at the bottom-left corner, over the sidebar's footer, where it
        # covers nothing that matters.
        lw = Inches(min(5.2, 0.5 + len(label) * 0.095))
        pill = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.25), H - Inches(0.62), lw, Inches(0.4))
        pill.adjustments[0] = 0.5
        pill.fill.solid()
        pill.fill.fore_color.rgb = NAVY
        pill.line.fill.background()
        pill.shadow.inherit = False
        text(s, Inches(0.45), H - Inches(0.62), lw - Inches(0.3), Inches(0.4), label, 12, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        s.notes_slide.notes_text_frame.text = label

    # Close
    s = prs.slides.add_slide(blank)
    fill(s, NAVY)
    circle(s, Inches(9.4), Inches(-2.2), Inches(6), 10)
    text(s, Inches(0.9), Inches(2.9), Inches(11.5), Inches(1.0), "Thank you", 44, bold=True)
    text(s, Inches(0.9), Inches(3.95), Inches(11.5), Inches(0.6), "controlytics.ai", 20, color=ACCENT)

    prs.save(str(OUT))
    print("wrote", OUT, len(SLIDES) + 2, "slides")


if __name__ == "__main__":
    main()
